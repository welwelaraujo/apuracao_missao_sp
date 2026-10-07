import json, os, pickle, collections
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter as L

S = os.path.dirname(os.path.abspath(__file__)); D = os.path.join(S, '..', 'tse'); OUTD = os.path.join(S, '..', 'dados')
A = pickle.load(open(os.path.join(D, 'agg.pkl'), 'rb'))
LOC = pickle.load(open(os.path.join(D, 'locais.pkl'), 'rb'))
valid, missao, names = A['valid'], A['missao'], A['names']
CARGO_NM = {'6': 'Deputado Federal', '7': 'Deputado Estadual', '3': 'Governador', '5': 'Senador', '1': 'Presidente'}

# status / nome de urna from TSE state result files
st = {}
for c in ('6', '7'):
    d = json.load(open(os.path.join(D, f'sp-c000{c}.json'), encoding='utf-8'))
    for a in d['carg'][0]['agr']:
        for p in a['par']:
            if p['n'] == '14':
                for x in p['cand']:
                    st[(c, x['n'])] = (x['nmu'], x['st'], int(x['vap']))
cargos = sorted({k[3] for k in missao}, key=lambda c: ['6', '7', '3', '5', '1'].index(c))
tot_c = collections.Counter(); tot_cn = collections.Counter()
for k, v in missao.items():
    tot_c[k[3]] += v; tot_cn[(k[3], k[4])] += v
valid_c = collections.Counter()
for k, v in valid.items():
    valid_c[k[3]] += v

cands = []
for c in cargos:
    for (cc, n), v in sorted(tot_cn.items(), key=lambda x: -x[1]):
        if cc != c: continue
        if len(n) == 2:
            nm, s = 'Voto de legenda (14)', 'Legenda'
        else:
            nm, s = (st.get((c, n), (names[(c, n)], '', 0)))[:2]
        cands.append(dict(c=c, n=n, nm=nm, st=s, v=v))
cidx = {(x['c'], x['n']): i for i, x in enumerate(cands)}

# municipality centroid fallback for locais without coordinates
mun_cfg = json.load(open(os.path.join(D, 'mun-e006257-cm.json'), encoding='utf-8'))
tse2ibge = {m['cd']: m['cdi'] for a in mun_cfg['abr'] if a['cd'].upper() == 'SP' for m in a['mu']}
geo = json.load(open(os.path.join(D, 'sp_mun.geojson'), encoding='utf-8'))
cent = {}
for ft in geo['features']:
    g = ft['geometry']; ring = g['coordinates'][0] if g['type'] == 'Polygon' else max((p[0] for p in g['coordinates']), key=len)
    cent[ft['properties']['codarea']] = (sum(p[1] for p in ring) / len(ring), sum(p[0] for p in ring) / len(ring))

exec(open(os.path.join(S,'pagedata.py'),encoding='utf-8').read())

# sanity vs TSE totals
for c in ('6', '7'):
    for (cc, n), (nm, s, vap) in st.items():
        if cc == c and tot_cn.get((c, n), 0) != vap:
            print('DIFF', c, n, nm, vap, tot_cn.get((c, n), 0))

# ---------------- Excel ----------------
OUT = os.path.join(OUTD, 'votacao_missao_sp_2026.xlsx')
F = 'Arial'; base = Font(name=F); bold = Font(name=F, bold=True); blue = Font(name=F, color='0000FF')
title = Font(name=F, bold=True, size=14); small = Font(name=F, size=9, italic=True, color='595959')
sect = Font(name=F, bold=True, size=12, color='1F4E78'); hf = Font(name=F, bold=True, color='FFFFFF')
hfill = PatternFill('solid', fgColor='1F4E78'); elfill = PatternFill('solid', fgColor='DDEBF7'); tfill = PatternFill('solid', fgColor='F2F2F2')
NUM = '#,##0;(#,##0);-'; PCT = '0.00%;(0.00%);-'
FONTE = 'Fonte: TSE – Dados abertos, votação por seção eleitoral SP 2026 (1º turno, 04/10/2026) e eleitorado por local de votação 2026.'


def hdr(ws, r, cols, height=32):
    for i, h in enumerate(cols):
        x = ws.cell(r, 1 + i, h); x.font = hf; x.fill = hfill
        x.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    ws.row_dimensions[r].height = height


def top(ws, t):
    ws['A1'] = t; ws['A1'].font = title; ws['A2'] = FONTE; ws['A2'].font = small
    ws.sheet_view.showGridLines = False


wb = Workbook(); ws = wb.active; ws.title = 'Resumo'
top(ws, 'Missão (14) – Votação em São Paulo, Eleições 2026 (1º turno)')
ws['A3'] = 'Números em azul = dados do TSE; em preto = fórmulas. Linhas em azul-claro = eleitos.'; ws['A3'].font = small
hdr(ws, 5, ['Cargo', 'Votos da Missão', 'Votos válidos em SP', '% dos válidos', 'Votos nominais', 'Votos de legenda'])
r = 6
for c in cargos:
    leg = tot_cn.get((c, '14'), 0)
    vals = [CARGO_NM[c], tot_c[c], valid_c[c], f'=IFERROR(B{r}/C{r},0)', f'=B{r}-F{r}', leg]
    for i, v in enumerate(vals):
        x = ws.cell(r, 1 + i, v); x.font = blue if i in (1, 2, 5) else base
        x.number_format = PCT if i == 3 else NUM
    r += 1
r += 2
for c in cargos:
    ws.cell(r, 1, f'{CARGO_NM[c]} – votos por candidato').font = sect; r += 1
    hdr(ws, r, ['Nº', 'Nome na urna', 'Situação', 'Votos em SP', '% dos votos da Missão', 'Municípios com voto', 'Locais com voto']); r += 1
    first = r
    tot_row = None
    for cd in [x for x in cands if x['c'] == c]:
        nmun = len({k[0] for k in missao if k[3] == c and k[4] == cd['n'] and missao[k] > 0})
        nloc = sum(1 for k in missao if k[3] == c and k[4] == cd['n'] and missao[k] > 0)
        vals = [int(cd['n']), cd['nm'], cd['st'], cd['v'], None, nmun, nloc]
        for i, v in enumerate(vals):
            x = ws.cell(r, 1 + i, v); x.font = blue if i in (0, 3, 5, 6) else base
            if i in (3, 5, 6): x.number_format = NUM
            if cd['st'].startswith('Eleito'): x.fill = elfill
        r += 1
    last = r - 1
    for rr in range(first, last + 1):
        x = ws.cell(rr, 5, f'=IFERROR(D{rr}/$D${r},0)'); x.number_format = PCT; x.font = base
    ws.cell(r, 2, 'Total').font = bold
    x = ws.cell(r, 4, f'=SUM(D{first}:D{last})'); x.font = bold; x.number_format = NUM
    for i in range(1, 8): ws.cell(r, i).fill = tfill
    r += 3
for col, w in zip('ABCDEFG', [20, 34, 18, 14, 14, 14, 14]): ws.column_dimensions[col].width = w

# Por município
ws = wb.create_sheet('Por município'); top(ws, 'Missão – votos por município')
cols = ['Município']
for c in cargos: cols += [f'{CARGO_NM[c]} – votos Missão', f'{CARGO_NM[c]} – válidos', f'{CARGO_NM[c]} – % Missão']
cols += ['Total Missão (todos os cargos)']
hdr(ws, 4, cols, 44)
mv = collections.Counter(); mval = collections.Counter()
for k, v in missao.items():
    mn = (LOC.get(k[:3]) or {}).get('mun') or A['locinfo'][k[:3]][0]; mv[(mn, k[3])] += v
for k, v in valid.items():
    mn = (LOC.get(k[:3]) or {}).get('mun') or A['locinfo'].get(k[:3], ('?',))[0]; mval[(mn, k[3])] += v
order = sorted(muns, key=lambda m: -sum(mv[(m, c)] for c in cargos))
for i, m in enumerate(order):
    rr = 5 + i; ws.cell(rr, 1, m).font = base; col = 2; sums = []
    for c in cargos:
        a = ws.cell(rr, col, mv[(m, c)]); a.font = blue; a.number_format = NUM
        b = ws.cell(rr, col + 1, mval[(m, c)]); b.font = blue; b.number_format = NUM
        p = ws.cell(rr, col + 2, f'=IFERROR({L(col)}{rr}/{L(col + 1)}{rr},0)'); p.number_format = PCT; p.font = base
        sums.append(f'{L(col)}{rr}'); col += 3
    t = ws.cell(rr, col, '=' + '+'.join(sums)); t.number_format = NUM; t.font = bold
lr = 4 + len(order); tr = lr + 1
ws.cell(tr, 1, 'Total SP').font = bold
for cc in range(2, 3 + 3 * len(cargos)):
    ws.cell(tr, cc).fill = tfill
    if (cc - 2) % 3 == 2 and cc < 2 + 3 * len(cargos):
        x = ws.cell(tr, cc, f'=IFERROR({L(cc - 2)}{tr}/{L(cc - 1)}{tr},0)'); x.number_format = PCT
    else:
        x = ws.cell(tr, cc, f'=SUM({L(cc)}5:{L(cc)}{lr})'); x.number_format = NUM
    x.font = bold
ws.column_dimensions['A'].width = 30
for cc in range(2, 3 + 3 * len(cargos)): ws.column_dimensions[L(cc)].width = 15
ws.freeze_panes = 'B5'; ws.auto_filter.ref = f'A4:{L(2 + 3 * len(cargos))}{lr}'

# Por local de votação
ws = wb.create_sheet('Por local de votação'); top(ws, 'Missão – votos por local de votação (escola)')
cols = ['Município', 'Zona', 'Local de votação', 'Bairro', 'Latitude', 'Longitude', 'Eleitores']
for c in cargos: cols += [f'{CARGO_NM[c]} – votos Missão', f'{CARGO_NM[c]} – válidos', f'{CARGO_NM[c]} – % Missão']
cols += ['Total Missão', 'Candidato da Missão mais votado no local']
hdr(ws, 4, cols, 44)
lv = collections.defaultdict(collections.Counter)
for k, v in missao.items(): lv[k[:3]][k[3]] += v
rows = []
for o in out_locs:
    pass
for k in locs:
    li = LOC.get(k) or dict(mun=A['locinfo'].get(k, ('?', '', ''))[0], nome=A['locinfo'].get(k, ('', 'Local ' + k[2], ''))[1], bairro='', lat=None, lon=None, eleitores=0)
    best = max(((cn, v) for cn, v in ((kk[3:], vv) for kk, vv in missao.items() if False)), default=None)
    rows.append((k, li))
best_loc = {}
for k, v in missao.items():
    if len(k[4]) > 2:
        b = best_loc.get(k[:3])
        if b is None or v > b[1]: best_loc[k[:3]] = ((k[3], k[4]), v)
rows.sort(key=lambda x: -sum(lv[x[0]].values()))
for i, (k, li) in enumerate(rows):
    rr = 5 + i
    vals = [li['mun'], int(k[1]), li['nome'], li['bairro'], li['lat'], li['lon'], li['eleitores']]
    for j, v in enumerate(vals):
        x = ws.cell(rr, 1 + j, v); x.font = base
        if j == 6: x.number_format = NUM
    col = 8; sums = []
    for c in cargos:
        a = ws.cell(rr, col, lv[k][c]); a.font = blue; a.number_format = NUM
        b = ws.cell(rr, col + 1, valid.get(k + (c,), 0)); b.font = blue; b.number_format = NUM
        p = ws.cell(rr, col + 2, f'=IFERROR({L(col)}{rr}/{L(col + 1)}{rr},0)'); p.number_format = PCT; p.font = base
        sums.append(f'{L(col)}{rr}'); col += 3
    t = ws.cell(rr, col, '=' + '+'.join(sums)); t.number_format = NUM; t.font = bold
    b = best_loc.get(k)
    if b:
        cd = cands[cidx[b[0]]]
        ws.cell(rr, col + 1, f"{cd['nm']} ({cd['n']}) – {b[1]}").font = base
for j, w in enumerate([24, 7, 44, 24, 11, 11, 11] + [14] * (3 * len(cargos)) + [12, 44]):
    ws.column_dimensions[L(1 + j)].width = w
ws.freeze_panes = 'D5'; ws.auto_filter.ref = f'A4:{L(9 + 3 * len(cargos))}{4 + len(rows)}'

# Possíveis dobradinhas
ws = wb.create_sheet('Possíveis dobradinhas'); top(ws, 'Missão – possíveis dobradinhas (votação parecida por local de votação)')
ws['A3'] = ('Correlação = Pearson entre o % dos válidos de cada candidato nos locais de votação de SP (locais com ≥30 votos válidos). '
            'Sobreposição = Σ min(fração dos votos de A no local, fração dos votos de B no local), de 0% a 100%. '
            'É um indício geográfico: o voto é secreto e não prova que foram os mesmos eleitores.')
ws['A3'].font = small
hdr(ws, 5, ['Candidato da Missão', 'Cargo', 'Votos em SP', 'Parceiro sugerido', 'Partido', 'Cargo do parceiro',
            'Votos do parceiro em SP', 'Correlação', 'Sobreposição'], 32)
rr = 6
for key, ps in sorted(PART.items(), key=lambda kv: -page_c[pidx[kv[0]]]['v']):
    a = page_c[pidx[key]]
    for oc, on, rc, ov in ps[:5]:
        b = page_c[pidx[(oc, on)]]
        vals = [f"{a['nm']} ({a['n']})", CARGO_NM[a['c']], a['v'], f"{b['nm']} ({b['n']})", b['par'], CARGO_NM[oc], b['v'], rc, ov]
        for j, v in enumerate(vals):
            x = ws.cell(rr, 1 + j, v); x.font = blue if j in (2, 6, 7, 8) else base
            if j in (2, 6): x.number_format = NUM
            if j == 7: x.number_format = '0.00'
            if j == 8: x.number_format = '0.0%'
        rr += 1
    rr += 1
for j, w in enumerate([34, 18, 12, 34, 16, 18, 14, 12, 13]): ws.column_dimensions[L(1 + j)].width = w
ws.freeze_panes = 'A6'

wb.save(OUT); print(OUT)
print({c: (tot_c[c], valid_c[c]) for c in cargos})
